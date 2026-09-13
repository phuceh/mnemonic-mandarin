import requests
import tempfile
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from supabase import create_client

SUPABASE_URL = "https://tfgdctbuhsoflecjymvw.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InRmZ2RjdGJ1aHNvZmxlY2p5bXZ3Iiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc3ODA1NzQ1NCwiZXhwIjoyMDkzNjMzNDU0fQ.Zg3BMGoq1rURF7VNZqUyuA-0Ys1258H9fctmtIdT7pc"
OUTPUT = "/workspaces/mnemonic-mandarin/HSK1_Memorize_Mandarin.docx"
hsk_level = 1  # Change this when generating HSK 2, 3, 4 etc

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

def set_paragraph_spacing(p, before=0, after=0):
    pPr = p._p.get_or_add_pPr()
    pPrTag = OxmlElement('w:spacing')
    pPrTag.set(qn('w:before'), str(before))
    pPrTag.set(qn('w:after'), str(after))
    pPr.append(pPrTag)

def set_font_color(run, hex_color):
    r = run._r
    rPr = r.get_or_add_rPr()
    color = OxmlElement('w:color')
    color.set(qn('w:val'), hex_color)
    rPr.append(color)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=80, after=80)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'E8DDD0')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def download_image(url):
    r = requests.get(url, timeout=10)
    tmp = tempfile.mktemp(suffix='.png')
    with open(tmp, 'wb') as f:
        f.write(r.content)
    return tmp

def add_word_entry(doc, word, word_number):
    # Word number + Chinese character — both centred on same line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=40, after=10)
    run_num = p.add_run(f'{word_number}.  ')
    run_num.font.size = Pt(9)
    set_font_color(run_num, 'b08060')
    run_char = p.add_run(word['chinese'])
    run_char.font.size = Pt(38)
    run_char.font.bold = True
    set_font_color(run_char, '2d1810')

    # Pinyin
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=8)
    run = p.add_run(word['pinyin'])
    run.font.size = Pt(14)
    set_font_color(run, '8b5a3a')

    # English
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=20)
    run = p.add_run(word['english'].upper())
    run.font.size = Pt(11)
    run.font.bold = True
    set_font_color(run, 'b08060')

    # MNEMONIC label
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4)
    run = p.add_run('MNEMONIC')
    run.font.size = Pt(8)
    run.font.bold = True
    set_font_color(run, 'c0392b')

    # Mnemonic text
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=0)
    run = p.add_run(word['mnemonic'])
    run.font.size = Pt(11)
    run.font.italic = True
    set_font_color(run, '2d1810')

    # Blank line between mnemonic and image
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=0)
    run = p.add_run('')
    run.font.size = Pt(8)

    # Image
    if word.get('image_url'):
        try:
            tmp = download_image(word['image_url'])
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0, after=10)
            run = p.add_run()
            run.add_picture(tmp, width=Inches(3.4))
            os.unlink(tmp)
        except Exception as e:
            print(f"  Image error for {word['chinese']}: {e}")

# ── Main ──────────────────────────────────────────────────────────────────────
print(f"Fetching HSK{hsk_level} words...")
response = supabase.table("vocabulary").select(
    "word_id, chinese, pinyin, english, mnemonic, image_url"
).eq("hsk_level", hsk_level).order("word_id").execute()
words = response.data
print(f"Found {len(words)} words")

doc = Document()

# Page setup
section = doc.sections[0]
section.page_width = Inches(6)
section.page_height = Inches(9)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)

# Default style
style = doc.styles['Normal']
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = 1.0

# ── Title page ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=1000, after=60)
run = p.add_run('Memorize Mandarin')
run.font.size = Pt(28)
run.font.bold = True
set_font_color(run, '2d1810')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=60)
run = p.add_run(f'HSK {hsk_level} Vocabulary')
run.font.size = Pt(20)
set_font_color(run, 'c0392b')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=60)
run = p.add_run('Learn Mandarin words that actually stick\nwith mnemonic stories')
run.font.size = Pt(12)
set_font_color(run, '8b5a3a')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=0)
run = p.add_run('J. Fisk')
run.font.size = Pt(12)
run.font.bold = True
set_font_color(run, '2d1810')

doc.add_page_break()

# ── Introduction page ─────────────────────────────────────────────────────────
p = doc.add_paragraph()
set_paragraph_spacing(p, before=200, after=60)
run = p.add_run('Introduction')
run.font.size = Pt(18)
run.font.bold = True
set_font_color(run, '2d1810')

p = doc.add_paragraph()
set_paragraph_spacing(p, before=0, after=40)
run = p.add_run(
    f'Welcome to Memorize Mandarin - HSK {hsk_level} Vocabulary.\n\n'
    'This book is designed to be used alongside your Mandarin studies as a powerful memorisation tool. '
    'Rather than relying on repetition alone, each word is paired with a vivid mnemonic story that '
    'phonetically encodes the sound. Your brain is wired to remember stories, images and emotions far '
    'more effectively than lists - and that is exactly what this book exploits.\n\n'
    'For each word, read the mnemonic out loud and picture the scene as vividly as possible. '
    'The more absurd and vivid the image in your mind, the more likely it is to stick. '
    'Work through the book at your own pace, and return to any words you find difficult. '
    'Over time you will find that the mnemonics come to mind automatically whenever you encounter '
    'the word - which is exactly the goal.\n\n'
    'Good luck with your Mandarin journey. \u52a0\u6cb9\uff01'
)
run.font.size = Pt(11)
set_font_color(run, '2d1810')

doc.add_page_break()

# ── Word pages - 2 per page ───────────────────────────────────────────────────
for i in range(0, len(words), 2):
    add_word_entry(doc, words[i], i + 1)
    add_horizontal_rule(doc)
    if i + 1 < len(words):
        add_word_entry(doc, words[i + 1], i + 2)
    doc.add_page_break()

# ── Final page ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=1000, after=40)
run = p.add_run('Want the full experience?')
run.font.size = Pt(18)
run.font.bold = True
set_font_color(run, '2d1810')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=40)
run = p.add_run(
    f'This book covers HSK {hsk_level}. The full Memorize Mandarin app includes multiple HSK levels '
    'with audio pronunciation, interactive flashcards, quizzes and progress tracking. '
    'New levels are added regularly.'
)
run.font.size = Pt(11)
set_font_color(run, '5a3a2a')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=40, after=0)
run = p.add_run('memorizemandarin.com')
run.font.size = Pt(16)
run.font.bold = True
set_font_color(run, 'c0392b')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=20, after=0)
run = p.add_run('Free trial - no account needed')
run.font.size = Pt(11)
set_font_color(run, '8b5a3a')

doc.save(OUTPUT)
print(f"Done! Saved to {OUTPUT}")